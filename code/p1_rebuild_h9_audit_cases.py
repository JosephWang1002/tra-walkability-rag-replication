# -*- coding: utf-8 -*-
"""Rebuild Appendix H9 Word with true high + large-gap Full RAG cases.

Also resolve each retrieved chunk back to its parent corpus document title.
"""
import json
import os
import re
import shutil
import glob
import sys

import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8")

OUT_DIR = r"C:\Users\cosgl\Desktop\TRA\TRA\audit case 4"
ABL = r"E:\PAPER1\revision\results\ablation_3step\full_rag_v2"
RERUN = os.path.join(OUT_DIR, "rerun_full_rag", "full_rag_v2")
HUMAN = r"E:\PAPER1\revision\data\p0_01_200img_joined.csv"
CORPUS = r"E:\PAPER1\revision\data\corpus_with_source.json"
IMG_DIRS = [
    r"E:\PAPER1\human_sampled_images_200",
    os.path.join(OUT_DIR, "images"),
    os.path.join(OUT_DIR, "rerun_images"),
    r"C:\Users\cosgl\Desktop\TRA\TRA\validation_set_images_200",
]
DIMS = [
    "justice",
    "safety",
    "livability",
    "sustainability",
    "aesthetics",
    "comfort",
    "walkability",
]
MARKERS = [
    "TWSI",
    "Walk Score",
    "WalkScore",
    "space syntax",
    "Appleyard",
    "practicability",
    "urbanity",
    "hospitability",
    "PEDSAFE",
    "Vision Zero",
    "Woonerf",
    "15-minute",
    "complete streets",
    "Level of Service",
    "suitability",
]

CASES = [
    (
        "33971",
        "Case 1 — High score",
        "Highest Full-RAG walkability on N=200 validation (new Appendix C)",
        ABL,
    ),
    (
        "15105",
        "Case 2 — Low score",
        "Low walkability exemplar (new Appendix C re-run)",
        RERUN,
    ),
    (
        "21279",
        "Case 3 — Large human–LLM gap",
        "Large |human-LLM| walkability gap under Full RAG on N=200 (LLM under-scores)",
        ABL,
    ),
    (
        "25575",
        "Case 4 — Hallucination audit",
        "RAG hallucination probe under new Appendix C Full RAG (re-run)",
        RERUN,
    ),
]

_SKIP_LINE = re.compile(
    r"(?i)^("
    r"contents lists available|sciencedirect|journal homepage|show more|outline share|"
    r"get rights|under a creative|highlights|article info|a r t i c l e|keywords|"
    r"abstract$|available online|www\.elsevier|doi:|https?://|©|copyright|"
    r"article reuse|sagepub|journals\.sagepub|vol\.|volume |pp\.|page \d|"
    r"received |accepted |published|online \d|corresponding author|e-?mail|"
    r"school of |department of |faculty of |university |institute |"
    r"original article|original research|research article|review article|"
    r"planning research|resume article info|article info$|"
    r"issn |f\d+$|research-article|\d{4}/\d{2}/\d{2}"
    r")"
)
_JOURNALISH = re.compile(
    r"(?i)^(journal of|cities\d|case studies on|annals of|"
    r"safety and security engineering|frontiers in)\b"
)
_AUTHORISH = re.compile(
    r"(?i)^([A-ZÁÉÍÓÚÄÖÜ][a-zà-ÿ'\-]+(?:\s+[A-ZÁÉÍÓÚÄÖÜ][a-zà-ÿ'\-]+){0,3}"
    r"(?:,|\s+and\s+|\s+[0-9a-d,*]+|\s*$))"
)


def resolve(oid, preferred):
    p = os.path.join(preferred, f"{oid}.json")
    if os.path.exists(p):
        return p
    for d in [ABL, RERUN]:
        p = os.path.join(d, f"{oid}.json")
        if os.path.exists(p):
            return p
    raise FileNotFoundError(oid)


def norm_ws(t: str) -> str:
    return re.sub(r"\s+", " ", t).strip()


def _clean_title_line(ln: str) -> str:
    ln = re.sub(r"\s+\d{4},\s*Vol\..*$", "", ln).strip()
    ln = re.sub(r"\s+\d{4}/\d{2}/\d{2}\s+\d{2}:\d{2}\s*", "", ln).strip(" .…")
    return ln


def _looks_author(ln: str) -> bool:
    if re.search(r"(?i)\bet al\.?\b", ln):
        return True
    # "J. Aaron Hipp1*" / "Adam Millard-Ball1 , Ben ..."
    if re.match(r"^[A-Z]\.\s+[A-Z]", ln):
        return True
    if re.search(r"[A-Z][A-Za-z'\-]+\d[\s,*]", ln) and ("," in ln or "*" in ln):
        return True
    # "Jiaxuan Ea, Bo Xiaa,*" / "Hannah Hooka,b,*"
    if re.search(r"[A-Z][a-zà-ÿ'\-]+[a-d](?:,|\*|\s)", ln) and re.search(
        r"[A-Z][a-zà-ÿ'\-]+\s+[A-Z]", ln
    ):
        return True
    # "Laura Martínez-Buelvas a , Andry ..."
    if re.search(
        r"[A-Z][a-zà-ÿ'\-]+(?:\s+[A-Z][a-zà-ÿ'\-]+)+\s+[a-d]\b",
        ln,
    ):
        return True
    # bare "Kaylyn Levine" / "Frauke Luise Berghoefer*, Mark Vollrath"
    if re.match(r"^[A-Z][a-zà-ÿ'\-]+(?:\s+[A-Z][a-zà-ÿ'\-]+){1,3}\*?$", ln) and len(ln) < 50:
        return True
    if "*" in ln and "," in ln and re.match(r"^[A-Z]", ln) and len(ln) < 120:
        if re.search(r"[A-Z][a-zà-ÿ'\-]+\*", ln):
            return True
    return False


def _join_title_parts(parts: list[str]) -> str:
    title = ""
    for ln in parts:
        if not ln:
            continue
        if not title:
            title = ln
            continue
        open_end = bool(
            re.search(
                r"(?i)(\b(and|the|of|in|for|to|a|an|with|on|from|by|into|through|vs|a)\b|[:\-–,])$",
                title.rstrip(),
            )
        )
        if open_end or len(title) < 95 or ln[:1].islower():
            title = title + " " + ln
        else:
            # Sage / Elsevier wrap: keep joining while still short
            if len(title) < 140 and not _looks_author(ln):
                title = title + " " + ln
            else:
                break
    title = norm_ws(title).strip(" .")
    title = re.sub(
        r"(?i)^(Research Article|Original Article|Original Research|Review Article)\s+",
        "",
        title,
    ).strip()
    return title[:260]


def extract_doc_title(text: str) -> str:
    """Best-effort article/policy title from the start of a corpus document."""
    raw = text.replace("\xa0", " ")
    lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    if not lines:
        return "(untitled)"

    if re.search(r"[\u4e00-\u9fff]", lines[0]) and len(lines[0]) >= 4:
        return lines[0][:200]

    # Prefer block after journal homepage / doi (multi-line titles).
    for marker in (
        r"(?i)^journal homepage:",
        r"(?i)^doi:\s*10\.",
        r"(?i)^Case Studies on Transport Policy$",
        r"(?i)^Journal of Planning Education and Research$",
    ):
        for i, ln in enumerate(lines[:55]):
            if not re.match(marker, ln):
                continue
            parts = []
            for ln2 in lines[i + 1 : i + 14]:
                if re.match(
                    r"(?i)^(A R T I C L E|ARTICLE INFO|Abstract|ABSTRACT|Keywords|"
                    r"Highlights|Show more|Get rights|Outline|introduction)\b",
                    ln2,
                ):
                    break
                if _SKIP_LINE.match(ln2) or re.search(r"https?://|doi\.org|Volume ", ln2):
                    continue
                if _looks_author(ln2):
                    break
                cv = _clean_title_line(ln2)
                if not cv or len(cv) < 8:
                    continue
                if re.search(r"hDttOpsI|JPEXXX|</italic>", cv):
                    continue
                if _JOURNALISH.match(cv) and len(cv) < 90:
                    continue
                parts.append(cv)
            # Marker blocks already stop at author/abstract; join all parts.
            tit = norm_ws(" ".join(parts)).strip(" .")
            tit = re.sub(
                r"(?i)^(Research Article|Original Article|Original Research|Review Article)\s+",
                "",
                tit,
            ).strip()
            # Drop OCR garbage tokens
            tit = re.sub(r"\bhDttOpsI::?/\s*/?\S*", "", tit)
            tit = norm_ws(tit).strip(" .")
            if 15 <= len(tit) <= 320:
                return tit[:260]
    cands = []
    for ln in lines[:70]:
        if re.match(r"(?i)^(abstract|keywords|introduction|highlights|resume)\b", ln):
            break
        if _SKIP_LINE.match(ln):
            continue
        if re.search(r"https?://|doi\.org|10\.\d{4,}", ln):
            continue
        if re.match(r"^[\d\W_]+$", ln):
            continue
        if "…" in ln or ln.rstrip().endswith("..."):
            continue
        ln2 = _clean_title_line(ln)
        if not ln2 or len(ln2) < 8:
            continue
        if _JOURNALISH.match(ln2) and len(ln2) < 90:
            continue
        if "ORIGINAL RESEARCH" in ln2.upper() and len(ln2) < 90:
            continue
        if re.match(r"(?i)^research article$", ln2):
            continue
        if _looks_author(ln2) and cands:
            break
        if ln2.count(",") >= 2 and len(ln2) < 110 and re.search(r"[A-Z][a-z]+ [A-Z]", ln2):
            if cands:
                break
            continue
        cands.append(ln2)
        if sum(len(x) for x in cands) > 300:
            break

    if not cands:
        return norm_ws(lines[0])[:200]
    tit = _join_title_parts(cands)
    return tit if tit else "(untitled)"


class CorpusTitleIndex:
    def __init__(self, path: str):
        docs = json.load(open(path, encoding="utf-8"))
        self._docs = []
        for i, d in enumerate(docs):
            text = d.get("text") or ""
            self._docs.append(
                {
                    "i": i,
                    "source": d.get("source"),
                    "text": text,
                    "norm": norm_ws(text),
                    "title": extract_doc_title(text),
                }
            )

    def title_for_passage(self, p: dict) -> str:
        if p.get("source") == "policy":
            return (
                p.get("en_name")
                or p.get("zh_name")
                or p.get("policy_key")
                or "(policy document)"
            )
        for key in ("en_name", "zh_name", "doc_title", "title", "file", "file_name"):
            if p.get(key):
                return str(p[key])
        nt = norm_ws(p.get("text") or "")
        if not nt:
            return "(untitled)"
        for ntok in (40, 25, 15):
            key = " ".join(nt.split()[:ntok])
            if len(key) < 20:
                continue
            for d in self._docs:
                if d["source"] != "academic":
                    continue
                if key in d["norm"]:
                    return d["title"]
        return "(document title not resolved)"


def expl_lines(expl):
    lines = []
    for d in DIMS:
        if d not in expl:
            continue
        val = expl[d]
        lines.append(
            f"{d}: "
            + (" | ".join(map(str, val)) if isinstance(val, list) else str(val))
        )
    return lines


def hall_flags(expl, passages):
    ptxt = "\n".join(str(p.get("text", "")) for p in passages).lower()
    flags = []
    for d, val in (expl or {}).items():
        for c in val if isinstance(val, list) else [val]:
            c = str(c)
            absent = [
                m
                for m in MARKERS
                if re.search(re.escape(m), c, flags=re.I)
                and m.lower() not in ptxt
            ]
            if absent:
                flags.append({"dimension": d, "absent": absent, "reason": c})
    return flags


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(11)


def main():
    h = pd.read_csv(HUMAN, encoding="utf-8-sig")
    h["oid"] = h["oid"].astype(str)
    img_out = os.path.join(OUT_DIR, "images")
    os.makedirs(img_out, exist_ok=True)
    title_idx = CorpusTitleIndex(CORPUS)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    doc.add_heading(
        "Appendix H9 — Four Audit Cases (new Appendix C Full RAG)", level=0
    )
    add_para(
        doc,
        "Protocol: C2 scene → hybrid retrieval (4 academic + 1 English policy) → C3 (GPT-5.4). "
        "High/gap cases from N=200 Full RAG; low-score OID re-run under the same protocol. "
        "Each retrieved chunk is labelled with its parent document title.",
    )

    meta = []
    for oid, title, desc, pref in CASES:
        path = resolve(oid, pref)
        rec = json.load(open(path, encoding="utf-8"))
        assert rec.get("status") == "complete", (oid, rec.get("status"))
        flat = rec["scores_flat"]
        expl = (rec.get("step3_scoring") or {}).get("parsed_explanations") or {}
        passages = (rec.get("retrieval") or {}).get("passages") or []
        qmode = (rec.get("retrieval") or {}).get("query_mode")
        q = (rec.get("retrieval") or {}).get("query")
        flags = hall_flags(expl, passages)

        hits = []
        for d in IMG_DIRS:
            if os.path.isdir(d):
                hits += glob.glob(os.path.join(d, oid + "_*.jpg"))
        if hits:
            dst = os.path.join(img_out, os.path.basename(hits[0]))
            try:
                if os.path.abspath(hits[0]) != os.path.abspath(dst):
                    shutil.copy2(hits[0], dst)
            except PermissionError:
                pass

        chunk_meta = []
        print(oid, "walk", flat.get("walkability"), "src", path)
        for i, p in enumerate(passages[:5], 1):
            doc_title = title_idx.title_for_passage(p)
            chunk_meta.append({"rank": i, "source": p.get("source"), "doc_title": doc_title})
            print(f"  {i}. [{p.get('source')}] {doc_title}")

        doc.add_heading(f"{title} — OID {oid}", level=1)
        add_para(doc, desc)
        add_para(doc, f"Image ID (OID): {oid}")
        add_para(doc, f"Source JSON: {path}")
        add_para(doc, f"Retrieval query mode: {qmode}")
        if q:
            add_para(doc, "Retrieval query: " + str(q)[:600])

        doc.add_heading("Scores", level=2)
        if oid in set(h.oid):
            hr = h[h.oid == oid].iloc[0]
            table = doc.add_table(rows=1 + len(DIMS), cols=3)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Dimension"
            table.rows[0].cells[1].text = "Human"
            table.rows[0].cells[2].text = "LLM (Full RAG)"
            for i, d in enumerate(DIMS, 1):
                table.rows[i].cells[0].text = d
                table.rows[i].cells[1].text = str(float(hr[f"human_{d}"]))
                table.rows[i].cells[2].text = str(flat.get(d))
            gap = abs(float(flat["walkability"]) - float(hr["human_walkability"]))
            add_para(doc, f"Absolute walkability gap |LLM-human| = {gap:.1f}")
        else:
            table = doc.add_table(rows=1 + len(DIMS), cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Dimension"
            table.rows[0].cells[1].text = "LLM (Full RAG)"
            for i, d in enumerate(DIMS, 1):
                table.rows[i].cells[0].text = d
                table.rows[i].cells[1].text = str(flat.get(d))
            add_para(doc, "No human score in N=200 joined table.")

        doc.add_heading("Rationale (per dimension)", level=2)
        for line in expl_lines(expl):
            add_para(doc, "• " + line)

        doc.add_heading(
            "Hallucination check (marker in rationale, absent from Top-5)", level=2
        )
        if flags:
            for f in flags:
                add_para(
                    doc,
                    f"• [{f['dimension']}] absent={f['absent']}: {f['reason']}",
                )
        else:
            add_para(doc, "No marker-based term-absence hallucination detected.")

        doc.add_heading("Retrieved Top-5 chunks", level=2)
        for i, p in enumerate(passages[:5], 1):
            doc_title = title_idx.title_for_passage(p)
            src = p.get("source") or ""
            tx = (p.get("text") or "").strip()
            if len(tx) > 1200:
                tx = tx[:1200].rstrip() + " …"
            add_para(doc, f"{i}. [{src}] {doc_title}", bold=True)
            add_para(doc, tx)
        doc.add_page_break()
        meta.append(
            {
                "oid": oid,
                "title": title,
                "walk": flat.get("walkability"),
                "n_hall": len(flags),
                "json": path,
                "chunks": chunk_meta,
            }
        )

    out = os.path.join(OUT_DIR, "Appendix_H9_four_audit_cases_NEW_C_rerun.docx")
    try:
        doc.save(out)
    except PermissionError:
        out = os.path.join(
            OUT_DIR, "Appendix_H9_four_audit_cases_NEW_C_rerun_v2.docx"
        )
        doc.save(out)
    json.dump(
        {"cases": meta, "docx": out},
        open(os.path.join(OUT_DIR, "cases_summary_new_c.json"), "w", encoding="utf-8"),
        ensure_ascii=False,
        indent=2,
    )
    print("Wrote", out)


if __name__ == "__main__":
    main()
